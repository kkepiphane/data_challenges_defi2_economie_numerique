"""Rapport PowerPoint (10 diapositives) — python-pptx.

Usage : python -m src.rapport
Tous les chiffres et graphiques sont recalculés depuis data_processed/ avec les filtres par défaut
du tableau de bord (établissements de dépôt et de crédit, statuts « en activité » + « non renseigné »).
Graphiques natifs PowerPoint (modifiables) ; cartes et matrice en images matplotlib.
"""
from __future__ import annotations

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE,
                            XL_TICK_LABEL_POSITION, XL_TICK_MARK)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import config as C
from . import figures as FIG
from . import indicators as I
from . import recommandations as R
from . import theme as TH
from .data import charger_tout, populations
from .formatage import compact, entier, km, nombre, pct, pp

SORTIE = C.OUTPUTS_DIR / "rapport_inclusion_numerique_togo.pptx"
POLICE = "Calibri"
W, H = 13.333, 7.5


def _t(x: str) -> str:
    """Espace fine insécable → insécable classique (mieux rendue par PowerPoint et WPS)."""
    return str(x).replace(" ", " ")


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


INK, INK2, MUTED = rgb(TH.INK), rgb(TH.INK_2), rgb(TH.MUTED)
VERT, VERT_F, OR, INDIGO, ROUGE = rgb(TH.GREEN_700), rgb(TH.VERT_FONCE), rgb(TH.GOLD), rgb(TH.INDIGO), rgb(TH.CRITIQUE)
BLANC, TEINTE, LIGNE = rgb("#FFFFFF"), rgb(TH.MENTHE), rgb(TH.LINE)
ARMOIRIES = C.ASSETS_DIR / "armoiries.png"


# ======================================================================================
# Primitives
# ======================================================================================
def texte(slide, x, y, w, h, paragraphes, taille=14, couleur=INK, gras=False, align=PP_ALIGN.LEFT,
          ancre=MSO_ANCHOR.TOP, interligne=1.08, espace_apres=0):
    """paragraphes : str, ou liste de paragraphes ; un paragraphe = str ou liste de (texte, {style})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancre
    if isinstance(paragraphes, str):
        paragraphes = [paragraphes]
    for i, par in enumerate(paragraphes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = interligne
        p.space_after = Pt(espace_apres)
        runs = [(par, {})] if isinstance(par, str) else par
        for t, st in runs:
            r = p.add_run()
            r.text = _t(t)
            f = r.font
            f.name = POLICE
            f.size = Pt(st.get("taille", taille))
            f.bold = st.get("gras", gras)
            f.italic = st.get("italique", False)
            f.color.rgb = st.get("couleur", couleur)
    return tb


def boite(slide, x, y, w, h, fond=TEINTE, rayon=0.08, ligne=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = rayon
    s.fill.solid()
    s.fill.fore_color.rgb = fond
    if ligne is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = ligne
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def pastille(slide, x, y, d, fond, libelle, couleur=BLANC, taille=12):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fond
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = _t(libelle)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = POLICE, Pt(taille), True, couleur
    return s


def bandeau(slide):
    """Filet vert en tête de diapositive (couverture et conclusion)."""
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(0.14))
    r.fill.solid()
    r.fill.fore_color.rgb = VERT_F
    r.line.fill.background()
    r.shadow.inherit = False


def fond(slide, couleur):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = couleur


def entete(slide, numero: int, rubrique: str, titre: str, message: str = ""):
    texte(slide, 0.6, 0.42, 9, 0.3, [[(f"{numero:02d}  ·  ", {"couleur": OR, "gras": True}),
                                      (rubrique.upper(), {"couleur": VERT, "gras": True})]], taille=11)
    texte(slide, 0.6, 0.72, 12.1, 0.9, titre, taille=28 if len(titre) <= 64 else 24, gras=True, couleur=INK, interligne=1.0)
    if message:
        texte(slide, 0.6, 1.55, 12.1, 0.5, message, taille=14, couleur=INK2)


def pied(slide, sources: str, numero: int, sombre: bool = False):
    c = rgb("#9FB8AA") if sombre else MUTED
    texte(slide, 0.6, 7.02, 11.2, 0.3, "Sources : " + sources, taille=9, couleur=c)
    texte(slide, 12.2, 7.02, 0.55, 0.3, str(numero), taille=9, couleur=c, align=PP_ALIGN.RIGHT)


def chiffre(slide, x, y, w, valeur, libelle, couleur=VERT, taille=34, taille_lib=11.5, couleur_lib=INK2):
    texte(slide, x, y, w, 0.62, valeur, taille=taille, gras=True, couleur=couleur, interligne=1.0)
    texte(slide, x, y + 0.66, w, 0.6, libelle, taille=taille_lib, couleur=couleur_lib)


def image(slide, chemin, x, y, w=None, h=None):
    if chemin is None or not chemin.exists():
        boite(slide, x, y, w or 3, h or 3, fond=TEINTE)
        texte(slide, x + 0.2, y + 0.2, (w or 3) - 0.4, 1, C.NON_DISPONIBLE, taille=11, couleur=MUTED)
        return
    slide.shapes.add_picture(str(chemin), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def _axe(ax, taille=10, format_nombre=None, visible=True, grille=True):
    ax.tick_labels.font.size = Pt(taille)
    ax.tick_labels.font.name = POLICE
    ax.tick_labels.font.color.rgb = MUTED
    if visible:
        ax.format.line.color.rgb = rgb(TH.AXIS)
    else:
        ax.format.line.fill.background()
    if format_nombre:
        ax.tick_labels.number_format = format_nombre
        ax.tick_labels.number_format_is_linked = False
    if hasattr(ax, "has_major_gridlines"):
        ax.has_major_gridlines = grille
        if grille:
            ax.major_gridlines.format.line.color.rgb = rgb(TH.GRID)
            ax.major_gridlines.format.line.width = Pt(0.75)


def graphe(slide, type_, x, y, w, h, categories, series: dict[str, list], couleurs: list[str], legende=True,
           format_valeurs=None, etiquettes=False, format_etiquettes="#,##0", taille=10):
    cd = CategoryChartData()
    cd.categories = categories
    for nom, vals in series.items():
        cd.add_series(nom, vals)
    gf = slide.shapes.add_chart(type_, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name, ch.font.size = POLICE, Pt(taille)
    ch.font.color.rgb = INK2
    ch.has_title = False
    ch.has_legend = legende and len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(taille)
        ch.legend.font.color.rgb = INK2
    plot = ch.plots[0]
    for s, coul in zip(plot.series, couleurs):
        if type_ in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
            s.format.line.color.rgb = rgb(coul)
            s.format.line.width = Pt(2.25)
            s.smooth = False
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 6
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = rgb(coul)
            s.marker.format.line.color.rgb = BLANC
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(coul)
            s.format.line.color.rgb = BLANC
            s.format.line.width = Pt(1)
    if etiquettes:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = format_etiquettes
        dl.number_format_is_linked = False
        dl.font.size = Pt(taille - 0.5)
        dl.font.color.rgb = INK2
        if type_ in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    _axe(ch.category_axis, taille, visible=True, grille=False)
    _axe(ch.value_axis, taille, format_valeurs, visible=False, grille=True)
    return ch


def masquer_axe(ax):
    ax.visible = False
    ax.has_major_gridlines = False
    ax.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    ax.major_tick_mark = XL_TICK_MARK.NONE


def etiquettes(ch, textes: list[str], position=XL_LABEL_POSITION.OUTSIDE_END, taille=10):
    serie = ch.plots[0].series[0]
    vals = [v for v in serie.values if v is not None]
    if vals:
        ch.value_axis.minimum_scale = 0
        ch.value_axis.maximum_scale = max(vals) * 1.22
    for i, txt in enumerate(textes):
        etiquette_point(serie, i, txt, taille=taille, couleur=INK2, position=position)


def etiquette_point(serie, idx: int, txt: str, taille=11, couleur=INK, position=XL_LABEL_POSITION.RIGHT):
    dl = serie.points[idx].data_label
    dl.has_text_frame = True
    dl.text_frame.text = _t(txt)
    dl.position = position
    for p in dl.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(taille), True, couleur, POLICE


def tableau(slide, x, y, w, h, entetes, lignes, largeurs, taille=10.5, taille_entete=10):
    t = slide.shapes.add_table(len(lignes) + 1, len(entetes), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for j, lw in enumerate(largeurs):
        t.columns[j].width = Inches(lw)
    for i, rangee in enumerate([entetes] + lignes):
        for j, val in enumerate(rangee):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = VERT_F if i == 0 else (BLANC if i % 2 else TEINTE)
            tf = c.text_frame
            tf.word_wrap = True
            runs = [(val, {})] if isinstance(val, str) else val
            p = tf.paragraphs[0]
            for txt, st in runs:
                r = p.add_run()
                r.text = _t(txt)
                r.font.name = POLICE
                r.font.size = Pt(taille_entete if i == 0 else st.get("taille", taille))
                r.font.bold = i == 0 or st.get("gras", False)
                r.font.color.rgb = BLANC if i == 0 else st.get("couleur", INK)
    return t


def notes(slide, txt: str):
    slide.notes_slide.notes_text_frame.text = txt


# ======================================================================================
# Calculs (identiques au tableau de bord, filtres par défaut)
# ======================================================================================
def calculs() -> dict:
    T = charger_tout()
    pop = populations(T)
    f = I.Filtres()
    ag = I.filtrer_agents(T["agents"], f)
    fi = I.filtrer_finance(T["finance"], f)
    total = int(T["pop_commune"].population.sum())
    nat = I.synthese_nationale(ag, fi, total)
    reg = I.table_territoriale("region", ag, fi, pop, f)
    pref = I.table_territoriale("prefecture", ag, fi, pop, f)
    com = I.table_territoriale("commune", ag, fi, pop, f)
    bm = I.phases_internet(T["internet_bm"])
    tel = T["telecom"]
    S = lambda k: I.serie(tel, k)
    hhi = I.hhi_parts(tel).set_index("annee")
    mm_seul = com[com.statut_couverture == I.STATUT_MM_SEUL].sort_values("population", ascending=False)
    bandes = I.bandes_distance(ag)
    cant = I.presence_canton(ag, fi, f)
    reco_c = R.generer(com, "commune", nat["hab_par_agent"])
    reco_p = R.generer(pref, "prefecture", nat["hab_par_agent"])
    ctrl = T["controles"]
    s = bm.set_index("annee").valeur_pct
    return dict(T=T, f=f, ag=ag, fi=fi, total=total, nat=nat, reg=reg, pref=pref, com=com, bm=bm, s=s, tel=tel, S=S, hhi=hhi,
                mm_seul=mm_seul, bandes=bandes, cant=cant, reco_c=reco_c, reco_p=reco_p, ctrl=ctrl,
                avant=(s[2017] - s[2013]) / 4, apres=(s[2022] - s[2017]) / 5,
                pic=bm.loc[bm.variation_pp.idxmax()])


# ======================================================================================
# Diapositives
# ======================================================================================
def construire() -> Path:
    K = calculs()
    C.FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    f_carte = FIG.carte_hab_etab(K["pref"], K["fi"], C.FIGURES_DIR / "carte_hab_par_etablissement.png")
    f_mm = FIG.carte_mm_seul(K["ag"], K["com"], K["fi"], I.STATUT_MM_SEUL, C.FIGURES_DIR / "carte_communes_mm_seul.png")
    f_mat = FIG.matrice(K["pref"], 12, C.FIGURES_DIR / "matrice_prefectures.png")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    vierge = prs.slide_layouts[6]
    nat, s, S, hhi = K["nat"], K["s"], K["S"], K["hhi"]
    mm = K["mm_seul"]
    part_mm = 100 * mm.population.sum() / K["total"]
    src_pts = "Géoportail PRISE 2021/22 (agents) · géoportail, extraction 01/2025 (établissements) · RGPH-5 2022"

    # ---------------------------------------------------------------- 1. Titre et objectif
    sl = prs.slides.add_slide(vierge)
    bandeau(sl)
    image(sl, ARMOIRIES, 0.8, 0.55, h=1.45)
    texte(sl, 1.95, 0.95, 8, 0.35, "DÉFI 2 · ÉCONOMIE NUMÉRIQUE", taille=12, gras=True, couleur=VERT_F)
    texte(sl, 1.95, 1.3, 8, 0.35, "Rapport d'analyse et tableau de bord interactif", taille=12, couleur=INK2)
    texte(sl, 0.8, 2.3, 11.5, 1.5, "Adoption du numérique et inclusion financière par le mobile money au Togo",
          taille=38, gras=True, couleur=INK, interligne=1.0)
    texte(sl, 0.8, 3.8, 10.5, 0.9, "Objectif : mesurer, à partir des seules données disponibles, où le réseau d'agents mobile money "
          "supplée l'absence d'offre financière formelle — et où agir en priorité.", taille=16, couleur=INK2)
    for i, (v, l, c) in enumerate([
        (f"{nombre(s[2022])} %", "de la population utilise Internet (2022)", VERT_F),
        (entier(nat["agents"]), "agents mobile money géolocalisés (2021/22)", VERT_F),
        (f"{len(mm)} communes", f"sans établissement de dépôt/crédit recensé ({entier(mm.population.sum())} hab.)", ROUGE),
    ]):
        x = 0.8 + i * 4.05
        boite(sl, x, 4.9, 3.8, 1.65, fond=TEINTE, rayon=0.06, ligne=LIGNE)
        chiffre(sl, x + 0.3, 5.08, 3.3, v, l, couleur=c, taille=30)
    pied(sl, "RGPH-5 2022 (INSEED) · Géoportail PRISE · Banque mondiale (WDI/UIT) · séries sectorielles télécoms 2013–2019. "
             "Aucune valeur imputée.", 1)
    notes(sl, "Les trois chiffres sont recalculés par src/rapport.py depuis data_processed/ : WDI 2022 ; nombre de lignes du fichier "
              "agents ; communes (unités RGPH-5) avec agents > 0 et établissements (banque, micro-finance, mutuelle ; statuts en activité "
              "ou non renseigné) = 0.")

    # ---------------------------------------------------------------- 2. Problématique
    sl = prs.slides.add_slide(vierge)
    entete(sl, 2, "Problématique et enjeux", "Un réseau de proximité dense, une offre financière formelle rare",
           "Le mobile money peut-il compenser l'absence de banques et d'IMF ? Où l'accès reste-t-il fragile ?")
    cartes = [
        ("1", VERT, "Adoption du numérique", f"{nombre(s[2022])} %", f"des Togolais utilisent Internet en 2022, contre {nombre(s[2013])} % en 2013 ; "
         f"{pct(100 * S('T abonnés Internet mobiles (Toutes technologies)')[2019] / S('T abonnés Internet Fixe et Mobile (Toutes technologies)')[2019], 0)} "
         "des abonnements Internet sont mobiles (2019)."),
        ("2", OR, "Réseau mobile money", entier(nat["agents"]), f"agents présents dans les {K['ag'].commune.nunique()} communes, soit "
         f"{entier(nat['hab_par_agent'])} habitants par agent."),
        ("3", ROUGE, "Offre financière formelle", entier(nat["etablissements"]), f"établissements de dépôt/crédit : "
         f"{entier(nat['hab_par_etab'])} habitants par établissement, aucun dans {len(mm)} communes."),
    ]
    for i, (n, coul, titre, val, desc) in enumerate(cartes):
        x = 0.6 + i * 4.1
        boite(sl, x, 2.35, 3.85, 3.35, fond=TEINTE)
        pastille(sl, x + 0.3, 2.6, 0.5, coul, n, taille=14)
        texte(sl, x + 0.95, 2.66, 2.7, 0.4, titre, taille=15, gras=True)
        texte(sl, x + 0.3, 3.35, 3.3, 0.8, val, taille=40, gras=True, couleur=coul if coul != OR else rgb("#8A5A0E"))
        texte(sl, x + 0.3, 4.25, 3.3, 1.35, desc, taille=13, couleur=INK2)
    boite(sl, 0.6, 5.95, 12.1, 0.8, fond=rgb("#FBE9E7"))
    texte(sl, 0.9, 6.07, 11.6, 0.6, [[("Enjeu : ", {"gras": True, "couleur": ROUGE}),
                                      (f"{nombre(nat['agents_par_etab'], 0)} agents MM pour un établissement financier. Là où aucun établissement "
                                       "n'est recensé, l'agent MM est le seul point d'accès physique à un service financier.", {})]],
          taille=14, ancre=MSO_ANCHOR.MIDDLE)
    pied(sl, "Banque mondiale (WDI/UIT) · " + src_pts, 2)
    notes(sl, "Habitants par agent = population RGPH-5 2022 ÷ agents ; habitants par établissement = population ÷ établissements de "
              "dépôt/crédit ; agents par établissement = agents ÷ établissements.")

    # ---------------------------------------------------------------- 3. Données et qualité
    sl = prs.slides.add_slide(vierge)
    entete(sl, 3, "Données, périmètre et règles de qualité", "Huit fichiers sources, zéro valeur imputée",
           "Seuls les fichiers de data/ et, faute de géométrie, les contours préfectoraux du Défi 1 sont mobilisés.")
    lignes = [
        ["Agents mobile money", "19 788 points GPS", "2021/22 (PRISE)", "Canton"],
        ["Établissements financiers", "738 points GPS", "Non datée (extr. 2025)", "Canton"],
        ["Population RGPH-5", "759 lignes (5 niveaux)", "2022", "Canton/quartier"],
        ["Internet — Banque mondiale", "% de la population", "1960–2022", "National"],
        ["Abonnés Internet, technologies", "25 indicateurs", "2013–2019", "National"],
        ["Marché télécom (CA, parts…)", "10 indicateurs", "2013–2019", "National"],
        ["Schémas des couches (2)", "46 et 107 champs décrits", "—", "—"],
    ]
    tableau(sl, 0.6, 2.3, 7.4, 3.9, ["Source", "Contenu", "Période", "Maille la plus fine"], lignes, [2.45, 2.05, 1.6, 1.3], taille=11)
    ctrl = K["ctrl"]
    stats = [(f"{len(ctrl)}", f"contrôles automatiques rejoués à chaque exécution ({(ctrl.resultat == 'OK').sum()} conformes, "
                               f"{(ctrl.resultat != 'OK').sum()} écarts documentés)"),
             ("98,4 %", "des agents (99,2 % des établissements) situés dans la préfecture qu'ils déclarent"),
             ("116", "unités communales appariées à la population (Danyi 1 + 2 fusionnées dans la source)"),
             ("7 / 46", "champs du schéma agents présents dans le fichier (13 / 107 pour les établissements)")]
    for i, (v, l) in enumerate(stats):
        y = 2.3 + i * 0.98
        texte(sl, 8.45, y, 1.45, 0.6, v, taille=26, gras=True, couleur=VERT)
        texte(sl, 9.95, y + 0.05, 2.8, 0.9, l, taille=11.5, couleur=INK2)
    boite(sl, 0.6, 6.35, 12.1, 0.52, fond=TEINTE)
    texte(sl, 0.85, 6.4, 11.7, 0.42, [[("Jointure impossible : ", {"gras": True}),
                                       ("population × canton (orthographes divergentes, quartiers du Grand Lomé) — ratios calculés aux "
                                        "seuls niveaux région, préfecture, commune. ARPU exclu (unité incohérente).", {})]],
          taille=12, couleur=INK2, ancre=MSO_ANCHOR.MIDDLE)
    pied(sl, "data/ (8 fichiers) · contours COD-AB OCHA via Défi 1 · journal : data_processed/controles_qualite.csv", 3)
    notes(sl, "Dictionnaire de données et journal de qualité générés automatiquement : data_dictionary.md, data_processed/journal_qualite.csv. "
              "La période PRISE 2021/22 provient de l'audit du catalogue du géoportail (Défi 1) ; les agents fournis y sont identiques.")

    # ---------------------------------------------------------------- 4. Internet
    sl = prs.slides.add_slide(vierge)
    entete(sl, 4, "Évolution de l'usage d'Internet",
           f"L'usage d'Internet s'accélère après 2017 : {pp(K['avant'])} par an, puis {pp(K['apres'])}",
           f"Individus utilisant Internet, % de la population — de {nombre(s[2013])} % (2013) à {nombre(s[2022])} % (2022).")
    d = K["bm"][K["bm"].annee >= 2005]
    ch = graphe(sl, XL_CHART_TYPE.LINE_MARKERS, 0.5, 2.15, 8.3, 4.7, [str(a) for a in d.annee],
                {"Individus utilisant Internet (%)": [round(v, 2) for v in d.valeur_pct]}, [TH.GREEN_700], legende=False,
                format_valeurs='0" %"')
    etiquette_point(ch.plots[0].series[0], len(d) - 1, f"{nombre(s[2022])} %", position=XL_LABEL_POSITION.LEFT)
    pic = K["pic"]
    chiffre(sl, 9.3, 2.3, 3.5, f"×{nombre(K['apres'] / K['avant'])}", "gain annuel moyen 2017–2022 rapporté à 2013–2017", couleur=VERT)
    chiffre(sl, 9.3, 3.75, 3.5, pp(pic.variation_pp), f"plus forte hausse annuelle ({int(pic.annee)})", couleur=rgb("#8A5A0E"))
    pen = S("Taux de pénétration Internet (Toutes technologies) (%)")
    chiffre(sl, 9.3, 5.2, 3.5, f"{nombre(pen[2019])} %", f"pénétration par abonnements en 2019, contre {nombre(s[2019])} % d'individus : "
            "multi-abonnements, mesures non substituables", couleur=INDIGO, taille_lib=11)
    pied(sl, "Banque mondiale, WDI (donnée UIT) ; séries sectorielles 2013–2019. Aucune valeur 2023. Lecture descriptive, sans causalité.", 4)
    notes(sl, "Variation = taux(année) − taux(année précédente). Accélération : variation > variation précédente + 0,1 pt. Aucune variable "
              "explicative n'est disponible : les phases ne sont pas interprétées causalement.")

    # ---------------------------------------------------------------- 5. Télécoms
    sl = prs.slides.add_slide(vierge)
    der = hhi.iloc[-1]
    entete(sl, 5, "Résultats sur le marché des télécommunications",
           f"Un duopole serré ; {pct(100 * S('T abonnés Internet mobiles (Haut débit)')[2019] / S('T abonnés Internet mobiles (Toutes technologies)')[2019], 0)} "
           "de l'Internet mobile en haut débit (2019)",
           f"Parts de marché en abonnés mobiles 2019 : Togocom {pct(der.Togocom)}, Moov {pct(der.Moov)} — HHI {entier(der.hhi)}.")
    ans = [str(a) for a in hhi.index]
    ch = graphe(sl, XL_CHART_TYPE.LINE_MARKERS, 0.5, 2.15, 5.9, 3.55, ans,
                {"Togocom": list(hhi.Togocom), "Moov": list(hhi.Moov)}, [TH.GREEN_700, TH.INDIGO], format_valeurs='0" %"')
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 40, 60
    texte(sl, 0.6, 5.72, 5.8, 0.3, "Parts de marché en abonnés mobiles (%)", taille=11, couleur=MUTED)
    tel = K["tel"]
    tm = tel[tel.famille == "Internet mobile — technologie"]
    cats = [f"{o} {a}" for o in ["Togocom", "Moov"] for a in (2013, 2019)]
    ser = {}
    for tech in ["2G (GPRS/EDGE)", "3G", "4G"]:
        v = []
        for o in ["Togocom", "Moov"]:
            for a in (2013, 2019):
                x = tm[(tm.groupe == o) & (tm.technologie == tech) & (tm.annee == a)].valeur
                v.append(float(x.iloc[0]) if len(x) else None)
        ser[tech] = v
    ch2 = graphe(sl, XL_CHART_TYPE.COLUMN_STACKED, 6.75, 2.15, 6.0, 3.55, cats, ser, TH.ORDINALE_VERTE, format_valeurs='0.0,," M"')
    ch2.plots[0].gap_width = 70
    texte(sl, 6.85, 5.72, 5.9, 0.3, "Abonnés Internet mobile par technologie (4G Moov 2019 absente de la source)", taille=11, couleur=MUTED)
    ca, inv, ftth = S("Chiffres d'Affaires"), S("Investissement"), S("FTTH")
    for i, (v, l) in enumerate([(f"{compact(ca[2019])} FCFA", "chiffre d'affaires 2019"),
                                (pct(100 * inv[2019] / ca[2019]), "taux d'investissement 2019"),
                                (entier(ftth[2019]), "abonnés fibre (FTTH) en 2019"),
                                (compact(S("T abonnés Internet Fixe et Mobile (Toutes technologies)")[2019]), "abonnements Internet 2019")]):
        x = 0.6 + i * 3.05
        texte(sl, x, 6.12, 2.9, 0.45, v, taille=20, gras=True, couleur=VERT)
        texte(sl, x, 6.55, 2.9, 0.35, l, taille=11, couleur=INK2)
    pied(sl, "séries sectorielles 2013–2019 (observationdata-mesqyx, -cxnvmoc). Togo Cellulaire → Togocom ; Atlantique Telecom → Moov.", 5)
    notes(sl, "HHI = somme des carrés des parts de marché ; avec deux opérateurs il ne peut descendre sous 5 000. Taux d'investissement = "
              "investissement ÷ CA. Périmètre du CA non précisé dans la source ; ARPU exclu (unité incohérente).")

    # ---------------------------------------------------------------- 6. Répartition territoriale
    sl = prs.slides.add_slide(vierge)
    reg = K["reg"].set_index("territoire").reindex(C.REGIONS)
    hi, lo = reg.hab_par_etab.idxmax(), reg.hab_par_etab.idxmin()
    kp = K["pref"].set_index("territoire")
    entete(sl, 6, "Répartition territoriale des services financiers",
           f"Offre formelle : 1 établissement pour {entier(reg.hab_par_etab[hi])} habitants en {hi}",
           f"Contre {entier(reg.hab_par_etab[lo])} en {lo}. La préfecture de Kpendjal ({entier(kp.population['Kpendjal'])} hab.) "
           "n'a aucun établissement de dépôt/crédit recensé.")
    image(sl, f_carte, 0.5, 2.15, h=4.75)
    texte(sl, 4.4, 2.1, 4.3, 0.3, "Habitants par établissement de dépôt/crédit", taille=12, gras=True)
    ch = graphe(sl, XL_CHART_TYPE.BAR_CLUSTERED, 4.3, 2.45, 4.4, 2.95, list(reg.index),
                {"Habitants par établissement": [round(v) for v in reg.hab_par_etab]}, [TH.TERRACOTTA], legende=False)
    ch.category_axis.reverse_order = True
    ch.plots[0].gap_width = 80
    masquer_axe(ch.value_axis)
    etiquettes(ch, [entier(v) for v in reg.hab_par_etab])
    part_etab = 100 * reg.etablissements["Maritime"] / reg.etablissements.sum()
    part_pop = 100 * reg.population["Maritime"] / reg.population.sum()
    boite(sl, 4.4, 5.55, 4.3, 1.25, fond=TEINTE)
    texte(sl, 4.6, 5.65, 3.9, 1.05, [[(f"{pct(part_etab, 0)} ", {"gras": True, "couleur": VERT, "taille": 20}),
                                      ("des établissements sont en région Maritime, ", {}),
                                      (f"qui compte {pct(part_pop, 0)} de la population.", {})]],
          taille=12.5, couleur=INK2, ancre=MSO_ANCHOR.MIDDLE)
    catv = K["fi"].categorie.value_counts()
    sam = K["fi"][K["fi"].jours_renseignes].ouvert_samedi.mean() * 100
    for i, (v, l) in enumerate([(entier(catv.get("Micro-finance", 0)), "points de micro-finance"),
                                (entier(catv.get("Banque", 0)), "agences bancaires"),
                                (entier(catv.get("Mutuelle", 0)), "mutuelles"),
                                (pct(sam, 0), "des établissements ouverts le samedi")]):
        y = 2.15 + i * 1.18
        texte(sl, 9.3, y, 3.4, 0.55, v, taille=28, gras=True, couleur=VERT)
        texte(sl, 9.3, y + 0.58, 3.4, 0.4, l, taille=12, couleur=INK2)
    pied(sl, src_pts + " · contours COD-AB OCHA (Défi 1).", 6)
    notes(sl, "Carte : habitants par établissement de dépôt/crédit par préfecture ; points noirs = établissements (coordonnées GPS). "
              "Dénominateur : population RGPH-5 2022 de la même préfecture ; Maritime inclut le Grand Lomé. Catégories et statuts : "
              "banques, micro-finance, mutuelles ; en activité ou statut non renseigné.")

    # ---------------------------------------------------------------- 7. Rôle du mobile money
    sl = prs.slides.add_slide(vierge)
    b = K["bandes"].set_index("bande")
    cant = K["cant"].statut_couverture.value_counts()
    entete(sl, 7, "Rôle du mobile money dans l'accès territorial",
           f"L'agent MM est le seul point de service financier pour {entier(mm.population.sum())} habitants",
           f"{len(mm)} communes ({pct(part_mm)} de la population) comptent des agents MM mais aucun établissement de dépôt/crédit recensé.")
    image(sl, f_mm, 0.5, 2.15, h=4.75)
    texte(sl, 4.4, 2.1, 4.5, 0.3, "Distance de chaque agent au premier établissement", taille=12, gras=True)
    ch = graphe(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, 4.3, 2.5, 4.6, 3.1, list(b.index),
                {"Part des agents (%)": [round(v, 1) for v in b.part_pct]}, [TH.TERRACOTTA], legende=False)
    for i, pt in enumerate(ch.plots[0].series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(TH.RAMPE_CHAUDE[3 + i])
    ch.plots[0].gap_width = 60
    masquer_axe(ch.value_axis)
    etiquettes(ch, [pct(v) for v in b.part_pct])
    texte(sl, 4.4, 5.7, 4.5, 0.5, f"Part des {entier(nat['agents'])} agents MM, distance à vol d'oiseau", taille=11, couleur=MUTED)
    stats = [(nombre(nat["agents_par_etab"], 0), "agents MM pour un établissement de dépôt/crédit", VERT),
             (pct(b.part_pct[["5–10 km", "10–20 km", "≥ 20 km"]].sum()),
              f"des agents à plus de {C.SEUIL_ELOIGNEMENT_KM} km du premier établissement", VERT),
             (f"{entier(cant.get(I.STATUT_MM_SEUL, 0))} / {entier(len(K['cant']))}", "cantons avec agents MM et sans établissement", ROUGE)]
    for i, (v, l, c) in enumerate(stats):
        y = 2.15 + i * 1.5
        texte(sl, 9.3, y, 3.4, 0.6, v, taille=30, gras=True, couleur=c)
        texte(sl, 9.3, y + 0.66, 3.4, 0.6, l, taille=12, couleur=INK2)
    pied(sl, src_pts + ". Distances : BallTree haversine, sans réseau routier ni établissements hors frontière.", 7)
    notes(sl, "Une commune « sans établissement » l'est dans les données fournies : l'absence réelle du service n'est pas démontrable. "
              "Cantons : comptages sans population (appariement non fiable).")

    # ---------------------------------------------------------------- 8. Territoires prioritaires
    sl = prs.slides.add_slide(vierge)
    top = K["pref"].sort_values("score", ascending=False)
    entete(sl, 8, "Territoires prioritaires et inégalités observées",
           f"{', '.join(top.territoire.iloc[:3])} et {top.territoire.iloc[3]} cumulent les déficits d'accès",
           "Rang centile de 4 indicateurs par préfecture (1 = moins bien desservie parmi les 39) ; texte = valeur observée.")
    image(sl, f_mat, 0.5, 2.15, w=7.7)
    top10 = mm.head(10)
    texte(sl, 8.65, 2.15, 4.1, 0.3, "Communes sans établissement les plus peuplées", taille=12, gras=True)
    ch = graphe(sl, XL_CHART_TYPE.BAR_CLUSTERED, 8.55, 2.5, 4.25, 3.95, [f"{r.territoire}" for r in top10.itertuples()],
                {"Population 2022": [int(v) for v in top10.population]}, [TH.CRITIQUE], legende=False, taille=10)
    ch.category_axis.reverse_order = True
    ch.plots[0].gap_width = 60
    masquer_axe(ch.value_axis)
    etiquettes(ch, [entier(v) for v in top10.population], taille=9.5)
    texte(sl, 8.65, 6.5, 4.1, 0.4, f"{len(mm)} communes au total, {entier(mm.population.sum())} habitants (RGPH-5 2022).",
          taille=11, couleur=MUTED)
    pied(sl, src_pts + ". Rang moyen : moyenne simple, sans pondération.", 8)
    notes(sl, "Indicateurs : habitants par agent, habitants par établissement (aucun = rang le plus défavorable), distance médiane "
              "agent → établissement, part des agents à plus de 5 km. Rang moyen = moyenne des 4 rangs centiles.")

    # ---------------------------------------------------------------- 9. Recommandations
    sl = prs.slides.add_slide(vierge)
    rp = K["reco_p"]
    rc = K["reco_c"]
    p1 = rc[rc.priorite == "Priorité 1"]
    entete(sl, 9, "Recommandations stratégiques ciblées",
           f"{len(p1)} communes en priorité 1 : des déficits cumulés, trois leviers",
           "Priorité = nombre de critères remplis parmi C1 aucun établissement · C2 peu d'agents · C3 distance > 10 km · C4 agents dispersés.")
    sel = pd.concat([rp[rp.territoire == "Kpendjal"].assign(territoire="Préfecture de Kpendjal"),
                     p1[p1.criteres.str.contains("C1")].head(2), p1[~p1.criteres.str.contains("C1")].head(1)]).head(4)
    lignes = []
    for r in sel.itertuples():
        ind = (f"{entier(r.population)} hab. · {entier(r.agents)} agents MM · {entier(r.etablissements)} établ. · "
               f"distance médiane {km(r.dist_mediane_km)}")
        lignes.append([[(r.territoire, {"gras": True})], r.constats_courts.replace(" | ", " · "), ind,
                       r.actions_courtes.split(" | ")[0], "Suivi : " + r.suivi])
    tableau(sl, 0.6, 2.2, 12.1, 3.7, ["Territoire", "Problème observé", "Indicateurs sources", "Action proposée", "Impact attendu"],
            lignes, [1.8, 2.55, 2.55, 3.25, 1.95], taille=10.5)
    boite(sl, 0.6, 6.05, 12.1, 0.82, fond=TEINTE)
    texte(sl, 0.85, 6.1, 11.7, 0.72, [[("Leviers transverses : ", {"gras": True, "couleur": VERT}),
                                       ("(1) adosser les agents MM à une IMF pour le dépôt et l'épargne ; (2) densifier les agents là où "
                                        "l'on dépasse 2× la référence nationale d'habitants par agent ; (3) services financiers numériques "
                                        "là où le premier établissement est à plus de 10 km. Aucun objectif chiffré.", {})]],
          taille=11.5, couleur=INK2, ancre=MSO_ANCHOR.MIDDLE)
    pied(sl, f"RGPH-5 2022 · PRISE 2021/22 · géoportail 01/2025. Liste complète ({len(rc)} communes, {len(rp)} préfectures) : "
             "tableau de bord, page Recommandations.", 9)
    notes(sl, "Règles : C1 établissements = 0 et agents > 0 ; C2 habitants/agent > 2 × référence nationale (409) ; C3 distance médiane "
              "> 10 km ; C4 établissements > 0 et ≥ 50 % des agents à > 5 km. Les impacts désignent l'indicateur à suivre ; "
              "aucun effet n'est chiffré.")

    # ---------------------------------------------------------------- 10. Limites et conclusion
    sl = prs.slides.add_slide(vierge)
    bandeau(sl)
    entete(sl, 10, "Limites des données et conclusion", "Le mobile money maille le pays ; l'offre formelle doit s'y adosser")
    p1c = K["reco_c"][K["reco_c"].priorite == "Priorité 1"]
    concl = [("1", f"L'usage d'Internet progresse vite ({nombre(s[2013])} % → {nombre(s[2022])} %), sur un marché mobile en duopole équilibré."),
             ("2", f"Le réseau d'agents MM ({entier(nat['agents'])}) est {nombre(nat['agents_par_etab'], 0)} fois plus dense que celui des "
                   f"établissements ; {len(mm)} communes n'ont que lui."),
             ("3", f"Les {len(p1c)} communes de priorité 1 se situent en " + ", ".join(
                 f"{r} ({n})" for r, n in p1c.region.value_counts().items()) + ".")]
    for i, (n, t) in enumerate(concl):
        y = 2.0 + i * 0.95
        pastille(sl, 0.8, y, 0.48, VERT_F, n, taille=13)
        texte(sl, 1.5, y + 0.02, 5.4, 0.9, t, taille=14, couleur=INK)
    boite(sl, 7.3, 1.95, 5.4, 4.7, fond=rgb("#FBF4E6"), rayon=0.05, ligne=rgb("#EEDDBA"))
    texte(sl, 7.6, 2.15, 4.9, 0.4, "Limites — données non disponibles", taille=15, gras=True, couleur=rgb("#8A5A0E"))
    lim = ["Aucune date de collecte individuelle : pas d'évolution des points de service.",
           "Ni transactions, ni comptes, ni genre des agents : l'usage réel du MM n'est pas mesuré.",
           "Séries télécoms arrêtées en 2019 ; Internet national uniquement (pas de ventilation territoriale).",
           "Absence dans les données ≠ absence réelle ; distances à vol d'oiseau.",
           "Aucune analyse causale possible (pas de variable explicative)."]
    texte(sl, 7.6, 2.7, 4.9, 3.8, [[("—  ", {"couleur": OR, "gras": True}), (l, {})] for l in lim], taille=12.5,
          couleur=INK, espace_apres=7)
    boite(sl, 0.8, 5.55, 6.1, 1.1, fond=TEINTE, rayon=0.06, ligne=LIGNE)
    texte(sl, 1.05, 5.68, 5.6, 0.9, [[("Tableau de bord interactif : ", {"gras": True, "couleur": VERT_F}),
                                     ("7 pages, filtres globaux (période, territoire, opérateur, catégorie, statut), cartes, "
                                      "téléchargements CSV et fiche méthodologique — streamlit run app.py", {})]], taille=12.5, couleur=INK2)
    pied(sl, "ensemble des fichiers de data/ ; méthode et contrôles : methodologie_et_limites.md, data_dictionary.md.", 10)
    notes(sl, "Conclusion fondée exclusivement sur les indicateurs calculés ; les limites reprennent l'onglet « Analyses impossibles » du tableau de bord.")

    C.OUTPUTS_DIR.mkdir(exist_ok=True)
    prs.save(SORTIE)
    return SORTIE


if __name__ == "__main__":
    print(construire())
